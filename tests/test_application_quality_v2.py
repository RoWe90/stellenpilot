"""Regression tests for factual curation, ATS reading order and quality reports."""
from __future__ import annotations

import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from stellenpilot.cv_policy import curate
from stellenpilot.document_quality import (
    inspect_documents, letter_quality_issues, polish_generated_german,
)
from stellenpilot.tailor import build_cv, build_letter, keyword_plan, tailor_job


def example_profile():
    return {
        "person": {"name": "Test Person", "location": "Berlin",
                   "headlines": ["Softwareentwickler Python"],
                   "summary": "Softwareentwicklung und Automatisierung."},
        "skills": {"AI Engineering & Agentic Coding": [
            "Python", "OpenCode", "Software-Automatisierung"],
            "DevOps": ["CI/CD", "Azure Pipelines"]},
        "experience": [
            {"title": "Engineer", "company": "Test GmbH", "dates": "2022 - heute",
             "bullets": [f"Entwicklung und Automatisierung mit Python in Anwendung {i}."
                         for i in range(8)]},
            {"title": "Software Developer", "company": "Another GmbH",
             "dates": "2020 - 2022",
             "bullets": [f"Python-Tool für Datenanalyse Aufgabe {i}."
                         for i in range(5)]},
        ],
        "education": [
            {"degree": "Studium Informatik", "school": "Hochschule", "dates": "2020"},
            {"degree": "Studium Mathematik (ohne Abschluss)", "school": "Universität", "dates": ""},
            {"degree": "Abitur", "school": "Schule", "dates": "2015"},
        ],
    }


def example_plan(profile):
    return {"headline_index": 0, "summary": profile["person"]["summary"],
            "skill_groups": list(profile["skills"]),
            "experience": [list(range(len(e["bullets"]))) for e in profile["experience"]],
            "projects": []}


class CurateTests(unittest.TestCase):
    def test_balanced_preserves_profile_and_chronology(self):
        p = example_profile()
        original = deepcopy(p)
        plan = example_plan(p)
        selected, notes = curate(p, plan)
        self.assertEqual(len(selected["experience"][0]), 6)
        self.assertEqual(len(selected["experience"][1]), 3)
        self.assertEqual(selected["experience"][0][0], 0)
        self.assertEqual(p, original)
        self.assertEqual(len(plan["experience"][0]), 8)
        self.assertTrue(any("nicht im Bewerbungs-CV" in n for n in notes))

    def test_full_preserves_everything(self):
        p = example_profile()
        result, notes = curate(p, example_plan(p), "full", "full")
        self.assertEqual(result["experience"][0], list(range(8)))
        self.assertEqual(result["education_indices"], [0, 1, 2])
        self.assertEqual(notes, [])

    def test_essential_education_is_opt_in(self):
        p = example_profile()
        result, notes = curate(p, example_plan(p), "concise", "essential")
        self.assertEqual(result["education_indices"], [0, 2])
        self.assertIn("Ausbildungsstationen ausgelassen", "\n".join(notes))

    def test_invalid_options_rejected(self):
        p = example_profile()
        with self.assertRaises(ValueError):
            curate(p, example_plan(p), "guess")
        with self.assertRaises(ValueError):
            curate(p, example_plan(p), "full", "omit_everything")


class DashboardSettingTests(unittest.TestCase):
    def test_cv_controls_accept_allowed_values(self):
        from stellenpilot.server import clean_settings
        got = clean_settings({"tailor": {
            "cv_style": "lebenslauf", "cv_focus": "concise",
            "cv_education": "essential",
        }})
        self.assertEqual(got["tailor"]["cv_focus"], "concise")
        self.assertEqual(got["tailor"]["cv_education"], "essential")

    def test_cv_controls_reject_unknown_values(self):
        from stellenpilot.server import clean_settings
        with self.assertRaisesRegex(ValueError, "unknown cv focus"):
            clean_settings({"tailor": {"cv_focus": "delete-all"}})


class RenderingTests(unittest.TestCase):
    def test_modern_skills_are_readable_without_tables(self):
        profile = example_profile()
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / "cv.docx"
            build_cv(profile, example_plan(profile), dest, style="lebenslauf")
            doc = Document(dest)
            self.assertFalse(doc.tables)
            paragraphs = [p.text for p in doc.paragraphs]
            self.assertTrue(any("AI Engineering & Agentic Coding" in p and
                                "OpenCode" in p for p in paragraphs))
            self.assertTrue(any("CI/CD" in p for p in paragraphs))
            self.assertTrue(any("Studium Mathematik (ohne Abschluss)" in p
                                for p in paragraphs))
            self.assertAlmostEqual(doc.sections[0].page_width.mm, 210, delta=0.6)

    def test_education_can_be_essential_without_changing_profile(self):
        p = example_profile()
        selected, _ = curate(p, example_plan(p), "balanced", "essential")
        with tempfile.TemporaryDirectory() as tmp:
            docx = Path(tmp) / "cv.docx"
            build_cv(p, selected, docx, style="lebenslauf")
            doc = Document(docx)
            text = "\n".join(x.text for x in doc.paragraphs)
            self.assertIn("Studium Informatik", text)
            self.assertNotIn("Studium Mathematik", text)
            self.assertEqual(len(p["education"]), 3)

    def test_qa_catches_missing_evidence_without_leaking_it(self):
        p = example_profile()
        with tempfile.TemporaryDirectory() as tmp:
            cv, letter = Path(tmp) / "cv.docx", Path(tmp) / "letter.docx"
            build_cv(p, example_plan(p), cv, style="lebenslauf")
            build_letter({"company": "Company", "title": "Engineer"}, p,
                         "Hier ist mein Anschreiben.", letter, "de")
            good = inspect_documents(cv, letter, expected_bullets=[p["experience"][0]["bullets"][0]],
                                     require_pdf=False)
            self.assertTrue(good["ok"], good)
            bad = inspect_documents(cv, letter, expected_bullets=["confidential fact xyz"],
                                    require_pdf=False)
            self.assertFalse(bad["ok"])
            self.assertIn("missing_source_bullets", [i["code"] for i in bad["issues"]])
            self.assertNotIn("confidential fact xyz", str(bad))

    def test_qa_checks_page_count_and_pdf_dimensions(self):
        p = example_profile()
        with tempfile.TemporaryDirectory() as tmp:
            cv, letter = Path(tmp) / "cv.docx", Path(tmp) / "letter.docx"
            build_cv(p, example_plan(p), cv, style="lebenslauf")
            build_letter({"company": "Company", "title": "Engineer"}, p,
                         "Bewerbung mit Python.", letter, "de")
            w = PdfWriter()
            for _ in range(3):
                w.add_blank_page(width=595.3, height=841.9)
            cvpdf = cv.with_suffix(".pdf")
            with cvpdf.open("wb") as f:
                w.write(f)
            w2 = PdfWriter()
            w2.add_blank_page(width=595.3, height=841.9)
            letterpdf = letter.with_suffix(".pdf")
            with letterpdf.open("wb") as f:
                w2.write(f)
            checked = inspect_documents(cv, letter, [cvpdf, letterpdf])
            self.assertEqual(checked["stats"]["cv_pages"], 3)
            self.assertTrue(any(x["code"] == "too_many_pages" for x in checked["issues"]))
            self.assertTrue(checked["ok"])  # sparse-page warnings are not false build failures

    def test_generation_documents_qa_and_selection_report(self):
        p = example_profile()
        job = {"id": "unit-1", "company": "Demo AG", "title": "Python Entwickler",
               "description": "Python, Azure Pipelines CI/CD", "match": 90}
        with tempfile.TemporaryDirectory() as tmp:
            res = tailor_job(job, p, Path(tmp), make_pdf=False,
                             cv_style="lebenslauf", cv_focus="concise",
                             cv_education="essential")
            report = (Path(res["folder"]) / "report.md").read_text()
            self.assertIn("## CV-Auswahl", report)
            self.assertIn("## Dokumentenprüfung", report)
            self.assertIn("Status: PASS", report)
            self.assertIn("nicht im Bewerbungs-CV", report)
            self.assertTrue(res["quality"]["ok"])
            cvtext = "\n".join(x.text for x in Document(res["files"][0]).paragraphs)
            self.assertNotIn(p["experience"][0]["bullets"][7], cvtext)
            self.assertEqual(len(p["experience"][0]["bullets"]), 8)


class LanguageTests(unittest.TestCase):
    def test_polishes_only_known_compounds(self):
        source = "Python Datenverarbeitung, ETL Prozessen und KI Agenten Lösung sowie CI CD."
        polished = polish_generated_german(source)
        self.assertIn("Python-Datenverarbeitung", polished)
        self.assertIn("ETL-Prozessen", polished)
        self.assertIn("KI-Agenten-Lösung", polished)
        self.assertIn("CI/CD", polished)

    def test_style_checker_rejects_generic_letter(self):
        self.assertTrue(letter_quality_issues("Hiermit bewerbe ich mich.", lang="de"))


if __name__ == "__main__":
    unittest.main()
