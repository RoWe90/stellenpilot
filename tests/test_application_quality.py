"""Quality guard and DIN A4 export regression tests."""
import tempfile
import unittest
from pathlib import Path

from docx import Document

from stellenpilot.tailor import (
    LETTER_SYSTEM, LETTER_SYSTEM_DE, ProfileQualityError,
    build_cv, build_letter, keyword_plan, profile_quality_issues,
    tailor_job, template_letter_de,
)


PROFILE = {
    "person": {
        "name": "Example Developer", "location": "Berlin", "email": "example@sample.de",
        "languages": ["Deutsch (C2)", "English (B2)"],
        "headlines": ["Software Developer | Python"],
        "summary": "Softwareentwickler mit Erfahrung in Python-Projekten und Entwicklungsautomatisierung.",
    },
    "skills": {"Programming": ["Python"], "DevOps": ["Azure", "CI/CD"]},
    "experience": [
        {"title": "Software Developer", "company": "Example GmbH",
         "dates": "2021 - heute",
         "bullets": ["Entwicklung von Python-Werkzeugen zur Datenaufbereitung.",
                     "Automatisierung von Entwicklungsabläufen über Azure Pipelines."]},
    ],
    "projects": [],
    "education": [{"degree": "B.Sc. Informatik", "school": "Example Hochschule", "dates": "2021"}],
}
JOB = {"title": "Software Developer", "company": "Beispiel AG",
       "description": "Entwicklung mit Python und CI/CD."}


class ProfileQualityTests(unittest.TestCase):
    def test_clean_profile_has_no_warning(self):
        self.assertEqual(profile_quality_issues(PROFILE), [])

    def test_imported_long_title_with_bullets_is_rejected(self):
        bad = {**PROFILE, "experience": [
            {**PROFILE["experience"][0], "title": "Softwareentwickler • Datenverarbeitung • Python-Skripte"}
        ]}
        self.assertIn("Titel", profile_quality_issues(bad)[0])
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ProfileQualityError, "Berufsstation 1"):
                tailor_job(JOB, bad, Path(directory), make_pdf=False)
            self.assertEqual(list(Path(directory).iterdir()), [])

    def test_detects_cv_import_fragments_from_realistic_pdf_layout(self):
        role = {
            "title": "Werkstudent",
            "company": "Example GmbH, City, Entwicklung einer Web-Applikation für eine "
                       "Lernplattform (Vue, Azure, Pipelines, CI/CD)",
            "bullets": ["Entwicklung einer Web-Applikation (Vue, Azure",
                        "Pipelines, CI/CD)"],
        }
        issues = profile_quality_issues({"experience": [role]})
        self.assertTrue(any("Arbeitgeber" in item for item in issues), issues)
        self.assertTrue(any("Klammerausdrücke" in item for item in issues), issues)

    def test_detects_split_non_degree_studies_without_guessing(self):
        profile = {"experience": [], "education": [
            {"degree": "Fotograf (kein Abschluss", "school": "Umorientierung auf Studium)"},
            {"degree": "Studium Angewandte Informatik", "school": "Hochschule",
             "dates": "03/2021"},
        ]}
        issues = profile_quality_issues(profile)
        self.assertEqual(len(issues), 1)
        self.assertIn("Ausbildung 1", issues[0])

    def test_fragments_with_embedded_pdf_bullet_marks_are_blocked(self):
        bad = {**PROFILE, "experience": [{
            "title": "Bachelorant", "company": "Example GmbH",
            "bullets": ["Abschlussarbeit •", "und Implementierung der Software"],
        }]}
        issues = profile_quality_issues(bad)
        self.assertTrue(any("PDF-Aufzählungszeichen" in issue for issue in issues), issues)

    def test_imported_duplicate_bullets_flagged_without_dropping_facts(self):
        role = dict(PROFILE["experience"][0])
        role["bullets"] = ["Python Tool entwickelt.", "  Python   Tool entwickelt. "]
        bad = {**PROFILE, "experience": [role]}
        issues = profile_quality_issues(bad)
        self.assertTrue(any("doppelte" in issue for issue in issues))
        self.assertEqual(len(role["bullets"]), 2)


class DocumentLayoutTests(unittest.TestCase):
    def test_german_cv_is_din_a4_with_preserved_bullets(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "cv.docx"
            build_cv(PROFILE, keyword_plan(JOB, PROFILE), target, style="lebenslauf")
            doc = Document(target)
            section = doc.sections[0]
            self.assertAlmostEqual(section.page_width.mm, 210, delta=0.6)
            self.assertAlmostEqual(section.page_height.mm, 297, delta=0.6)
            self.assertTrue(any("BERUFSERFAHRUNG" in p.text for p in doc.paragraphs))
            self.assertTrue(any("Automatisierung von Entwicklungsabläufen" in p.text
                                for p in doc.paragraphs))
            header = next(p for p in doc.paragraphs
                          if p.text.split("\t")[0] == "Software Developer")
            self.assertTrue(header.paragraph_format.keep_with_next)
            self.assertTrue(any(p.text == "Example GmbH" for p in doc.paragraphs))
            self.assertIn("2021 - heute", header.text)
            self.assertTrue(any(PROFILE["experience"][0]["bullets"][0] in p.text
                                for p in doc.paragraphs))

    def test_letter_is_din_a4_and_has_distinct_subject(self):
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / "letter.docx"
            build_letter(JOB, PROFILE, "Mein Python-Projekt passt zur Stelle.", target, "de")
            doc = Document(target)
            self.assertAlmostEqual(doc.sections[0].page_width.mm, 210, delta=0.6)
            self.assertTrue(any(p.text.startswith("Bewerbung als") for p in doc.paragraphs))
            self.assertTrue(any("Mein Python-Projekt" in p.text for p in doc.paragraphs))


    def test_cv_and_letter_share_visual_identity(self):
        with tempfile.TemporaryDirectory() as directory:
            cv = Path(directory) / "cv.docx"
            letter = Path(directory) / "letter.docx"
            build_cv(PROFILE, keyword_plan(JOB, PROFILE), cv, style="lebenslauf")
            build_letter(JOB, PROFILE, "Python ist mein Schwerpunkt.", letter, lang="de")
            left, right = Document(cv), Document(letter)
            self.assertEqual(left.styles["Normal"].font.name, "Carlito")
            self.assertEqual(left.styles["Normal"].font.name, right.styles["Normal"].font.name)
            self.assertEqual(left.paragraphs[0].text, right.paragraphs[0].text)
            self.assertIn("w:pBdr", left.paragraphs[2]._p.xml)
            self.assertIn("w:pBdr", right.paragraphs[1]._p.xml)
            self.assertEqual(len(left.inline_shapes), 0)
            self.assertEqual(len(right.inline_shapes), 0)


class LetterPromptTests(unittest.TestCase):
    def test_prompts_discourage_unsolicited_deficits_and_inventions(self):
        self.assertIn("keine vermeintlichen Defizite", LETTER_SYSTEM_DE)
        self.assertIn("do not volunteer missing skills", LETTER_SYSTEM.lower())
        self.assertIn("Keine erfundenen", LETTER_SYSTEM_DE)

    def test_fallback_does_not_invent_missing_experience(self):
        plan = keyword_plan(JOB, PROFILE)
        plan.update({"covered": ["Python"], "learning": [], "missing": ["Kubernetes", "Terraform"]})
        text = template_letter_de(JOB, plan)
        self.assertNotIn("Kubernetes", text)
        self.assertNotIn("Terraform", text)
        self.assertNotIn("noch nicht", text)


if __name__ == "__main__":
    unittest.main()
