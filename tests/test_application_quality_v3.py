"""Regression tests for two-page evidence-aware balancing and unsupported job claims."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest import mock

from docx import Document
from pypdf import PdfWriter

from stellenpilot.cv_policy import (
    curate, expand_ranked_bullets, omitted_bullet_count, selection_notes,
)
from stellenpilot.document_quality import unsupported_action_claims
from stellenpilot.tailor import tailor_job


def profile():
    return {
        "person": {"name": "Example Applicant", "headlines": ["Developer Python"],
                   "summary": "Softwareentwickler mit Praxis in Python und ETL."},
        "skills": {"Programmiersprachen": ["Python"],
                   "DevOps": ["Azure Pipelines", "Docker"]},
        "experience": [
            {"title": "Developer", "company": "Example GmbH", "dates": "2022 – heute",
             "bullets": [f"Entwicklung und Automatisierung von Python-Datenpipelines, Thema {i}."
                         for i in range(8)]},
            {"title": "Junior Developer", "company": "Example GmbH", "dates": "2020 – 2021",
             "bullets": [f"Python-Werkzeug zur Datenaufbereitung, Teil {i}."
                         for i in range(5)]},
        ],
        "education": [{"degree": "B.Sc. Informatik", "school": "Hochschule", "dates": "2020"}],
    }


def plan(p):
    return {
        "headline_index": 0,
        "summary": p["person"]["summary"],
        "experience": [list(range(len(e["bullets"]))) for e in p["experience"]],
        "skill_groups": ["Programmiersprachen", "DevOps"],
        "projects": [],
        "coverage": 80,
    }


class SemanticGuardTests(unittest.TestCase):
    def test_unverified_acceptance_criteria_rejected(self):
        facts = "Python-ETL, Datenvalidierung, GitHub Copilot und OpenCode."
        draft = ("Anforderungen in prüfbare Akzeptanzkriterien zu überführen, "
                 "fällt mir durch meine ETL-Arbeit leicht.")
        self.assertIn("unsupported_experience_acceptance_criteria",
                      unsupported_action_claims(draft, facts))

    def test_deployment_cannot_be_inferred_from_tool_presence(self):
        facts = "Entwicklung von Azure Pipelines CI/CD und Einsatz von Docker."
        draft = "Für die Auslieferung nutze ich Azure Pipelines und Docker."
        self.assertIn("unsupported_experience_deployment",
                      unsupported_action_claims(draft, facts))

    def test_evidence_and_prospective_statements_are_allowed(self):
        facts = "Definition prüfbarer Akzeptanzkriterien und produktive Auslieferung mit Docker."
        draft = ("Ich definiere Akzeptanzkriterien für die Entwicklungsprojekte. "
                 "Ich nutze Docker für die Auslieferung der Software.")
        self.assertEqual(unsupported_action_claims(draft, facts), [])
        future = "Ich möchte künftig an der Definition von Akzeptanzkriterien mitwirken."
        self.assertEqual(unsupported_action_claims(future, ""), [])

    def test_source_must_support_specific_action_tool_pair(self):
        facts = ("CI/CD und Docker\n"
                 "Produktive Auslieferung einer Anwendung mit einer anderen Plattform")
        draft = "Für die Auslieferung nutze ich Azure Pipelines und Docker."
        self.assertIn("unsupported_experience_deployment",
                      unsupported_action_claims(draft, facts))

    def test_supported_technology_not_equivalent_to_supported_action(self):
        from stellenpilot.tailor import letter_body
        p = profile()
        p["skills"]["Programmiersprachen"].append("Pytest")
        fake = mock.Mock(enabled=True, max_chars=1500)
        fake.complete_json.side_effect = [
            {"body": "Für die Auslieferung nutze ich Azure Pipelines und Docker."},
            {"body": "Ich entwickle mit Python Datenpipelines und automatisiere "
                     "Verarbeitungsschritte.\n\n"
                     "In meiner beruflichen Arbeit entwickle ich dafür Python-Werkzeuge "
                     "zur Datenaufbereitung und Validierung.\n\n"
                     "Diese Erfahrung möchte ich in Ihr Entwicklungsteam einbringen."},
        ]
        text, note = letter_body(
            {"title": "Python Developer", "company": "Example AG",
             "description": "Python und Datenverarbeitung"}, p, plan(p), fake, "de")
        self.assertEqual(fake.complete_json.call_count, 2)
        self.assertNotIn("Auslieferung", text)
        self.assertNotIn("template used", note)


class PageBalanceTests(unittest.TestCase):
    def test_expansion_only_adds_documented_source_bullets(self):
        p = profile()
        full = plan(p)
        balanced, _ = curate(p, full, "balanced", "compact")
        self.assertEqual(omitted_bullet_count(p, balanced, full), 4)
        expanded = expand_ranked_bullets(p, balanced, full, 2)
        self.assertEqual(expanded["experience"][0], list(range(8)))
        self.assertEqual(expanded["experience"][1], list(range(3)))
        self.assertEqual(len(selection_notes(p, expanded)), 1)
        self.assertEqual(len(balanced["experience"][0]), 6)
        self.assertEqual(len(p["experience"][0]["bullets"]), 8)

    @staticmethod
    def fake_pdf_convert(source_paths, *, max_bullets: int):
        result = []
        for source in source_paths:
            source = Path(source)
            writer = PdfWriter()
            doc = Document(source)
            text = "\n".join(p.text for p in doc.paragraphs)
            count = sum("Python-Datenpipelines, Thema" in p.text or
                        "Python-Werkzeug zur Datenaufbereitung, Teil" in p.text
                        for p in doc.paragraphs)
            page_count = (3 if count > max_bullets else 2) if source.stem.endswith("_CV") else 1
            for _ in range(page_count):
                writer.add_blank_page(width=595.3, height=841.9)
            dest = source.with_suffix(".pdf")
            with dest.open("wb") as f:
                writer.write(f)
            result.append(dest)
        return result

    def _generate(self, limit):
        p = profile()
        job = {"title": "Python Developer", "company": "Example AG",
               "description": "Python, Datenpipelines und CI/CD", "match": 85}
        with tempfile.TemporaryDirectory() as tmp:
            def fake_convert(paths):
                return self.fake_pdf_convert(paths, max_bullets=limit)
            with mock.patch("stellenpilot.tailor.to_pdf", side_effect=fake_convert):
                res = tailor_job(job, p, Path(tmp), llm=None, make_pdf=True,
                                 cv_style="lebenslauf", cv_focus="balanced")
            cv = Document(next(Path(f) for f in res["files"] if f.endswith(".docx")
                               and Path(f).stem.endswith("_CV")))
            paragraphs = [p.text for p in cv.paragraphs]
            count = sum("Python-Datenpipelines, Thema" in x or
                        "Python-Werkzeug zur Datenaufbereitung, Teil" in x
                        for x in paragraphs)
            report = (Path(res["folder"]) / "edit" / "report.md").read_text()
            self.assertEqual(res["quality"]["stats"]["cv_pages"], 2)
            self.assertIn("## Dokumentenprüfung", report)
            return count, report

    def test_balanced_readds_everything_that_fits(self):
        count, report = self._generate(limit=13)
        self.assertEqual(count, 13)
        self.assertIn("4 zusätzliche belegte Tätigkeiten", report)

    def test_balanced_readds_only_as_many_as_fit(self):
        count, report = self._generate(limit=10)
        self.assertEqual(count, 10)
        self.assertIn("1 zusätzliche belegte Tätigkeit", report)


if __name__ == "__main__":
    unittest.main()
