"""Understated, ATS-readable layouts for CVs and cover letters.

The content is taken only from the validated profile, tailoring plan and letter.
All visual design is implemented with standard DOCX elements, without photos,
icons or decorative text columns.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

NAVY = RGBColor(0x1F, 0x35, 0x47)
TEAL = RGBColor(0x2B, 0x75, 0x82)
INK = RGBColor(0x29, 0x36, 0x43)
MUTED = RGBColor(0x64, 0x76, 0x82)
FONT = "Carlito"  # metrically compatible font installed in the Unraid image
HEADINGS = {
    "summary": "PROFIL", "skills": "KOMPETENZEN",
    "experience": "BERUFSERFAHRUNG", "projects": "PROJEKTE",
    "education": "AUSBILDUNG", "languages": "SPRACHEN",
}


def _font(run, *, size: float | None = None, bold: bool = False, color=None):
    run.bold = bold
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return run


def _rule(paragraph, *, color: str = "D9E3E7", width: int = 7, space: int = 3):
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, val in (("val", "single"), ("sz", str(width)), ("space", str(space)), ("color", color)):
        bottom.set(qn("w:" + key), val)
    borders.append(bottom)
    properties.append(borders)


def _no_split(row):
    row_properties = row._tr.get_or_add_trPr()
    row_properties.append(OxmlElement("w:cantSplit"))


def _cell_padding(cell, *, top=0, start=0, bottom=0, end=0):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        elt = OxmlElement(f"w:{edge}")
        elt.set(qn("w:w"), str(val))
        elt.set(qn("w:type"), "dxa")
        mar.append(elt)
    tc_pr.append(mar)


def _doc(*, kind: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(20)
    section.top_margin = Mm(19 if kind == "cv" else 23)
    section.bottom_margin = Mm(19 if kind == "cv" else 22)
    section.header_distance, section.footer_distance = Mm(9), Mm(11)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9.7 if kind == "cv" else 10.8)
    normal.font.color.rgb = INK
    normal.paragraph_format.line_spacing = 1.12 if kind == "cv" else 1.16
    normal.paragraph_format.space_after = Pt(3)
    doc.core_properties.title = "Bewerbungsunterlagen"
    doc.core_properties.subject = "Lebenslauf" if kind == "cv" else "Anschreiben"
    return doc


def _identity(doc, person: dict, *, subtitle: str = "", cv: bool = False):
    name = doc.add_paragraph()
    name.paragraph_format.space_after = Pt(2)
    _font(name.add_run(person.get("name") or ""), size=23 if cv else 21, bold=True, color=NAVY)
    if subtitle:
        role = doc.add_paragraph()
        role.paragraph_format.space_after = Pt(5)
        _font(role.add_run(subtitle), size=10.8, color=TEAL)
    contact = "  |  ".join(
        v for v in (person.get("location"), person.get("email"), person.get("phone"),
                    *((person.get("links") or [])[:2])) if v
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10 if cv else 18)
    _font(p.add_run(contact), size=9.1, color=MUTED)
    _rule(p, color="2B7582", width=12, space=8)


def _section(doc, label):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    _font(p.add_run(label), size=9.5, color=TEAL, bold=True)
    _rule(p, color="DDE6E9", width=5, space=5)


def _table(doc, widths):
    table = doc.add_table(rows=0, cols=len(widths))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Mm(width)
    return table


def _text_cell(cell, text, *, size=9.5, color=INK, bold=False, align=None):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _cell_padding(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.09
    if align is not None:
        p.alignment = align
    _font(p.add_run(str(text or "")), size=size, bold=bold, color=color)
    return p


def _experience(doc, entry: dict, indices):
    # One text paragraph with a right tab keeps the reading order accessible.
    # Its keep-with-next flag prevents stranded job headings at page boundaries.
    heading = doc.add_paragraph()
    heading.paragraph_format.tab_stops.add_tab_stop(Mm(169), WD_TAB_ALIGNMENT.RIGHT)
    heading.paragraph_format.keep_together = True
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(5)
    heading.paragraph_format.space_after = Pt(1)
    _font(heading.add_run(entry.get("title", "")), bold=True, size=10.5, color=NAVY)
    if entry.get("dates"):
        _font(heading.add_run("\t" + str(entry["dates"])), size=9.1, color=MUTED)

    employer = doc.add_paragraph()
    employer.paragraph_format.space_before = Pt(0)
    employer.paragraph_format.space_after = Pt(3)
    employer.paragraph_format.keep_with_next = bool(indices)
    _font(employer.add_run(entry.get("company", "")), size=9.3, color=MUTED)

    bullets = entry.get("bullets") or []
    valid = [j for j in indices if isinstance(j, int) and 0 <= j < len(bullets)]
    for index, j in enumerate(valid):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Mm(4)
        p.paragraph_format.first_line_indent = Mm(-2.6)
        p.paragraph_format.space_after = Pt(2.2)
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = index == 0 and len(valid) > 1
        _font(p.add_run("•  "), size=9.2, color=TEAL)
        _font(p.add_run(bullets[j]), size=9.35, color=INK)
    if valid:
        doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def build_cv(profile: dict, plan: dict, path: Path):
    doc = _doc(kind="cv")
    person = profile.get("person") or {}
    headlines = person.get("headlines") or []
    idx = plan.get("headline_index", 0)
    subtitle = headlines[idx] if isinstance(idx, int) and 0 <= idx < len(headlines) else (
        headlines[0] if headlines else "")
    _identity(doc, person, subtitle=subtitle, cv=True)

    if plan.get("summary"):
        _section(doc, HEADINGS["summary"])
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        _font(p.add_run(plan["summary"]), size=9.8, color=INK)

    skills = profile.get("skills") or {}
    if skills:
        _section(doc, HEADINGS["skills"])
        table = _table(doc, [47, 123])
        for group in plan.get("skill_groups", list(skills)):
            if group not in skills:
                continue
            row = table.add_row()
            _no_split(row)
            left, right = row.cells
            _text_cell(left, group, size=9.1, bold=True, color=NAVY)
            _text_cell(right, ", ".join(skills[group]), size=9.3)
            left.width, right.width = Mm(47), Mm(123)
            left.paragraphs[0].paragraph_format.space_after = Pt(2)
            right.paragraphs[0].paragraph_format.space_after = Pt(2)

    experience = profile.get("experience") or []
    if experience:
        _section(doc, HEADINGS["experience"])
        for entry, selected in zip(experience, plan.get("experience") or []):
            _experience(doc, entry, selected)

    projects = profile.get("projects") or []
    if projects:
        _section(doc, HEADINGS["projects"])
        for i in plan.get("projects") or []:
            if not isinstance(i, int) or not 0 <= i < len(projects):
                continue
            project = projects[i]
            _experience(doc, {"title": project.get("name", ""), "company": project.get("link", ""),
                              "bullets": project.get("bullets") or []},
                        list(range(len(project.get("bullets") or []))))

    education = profile.get("education") or []
    if education:
        _section(doc, HEADINGS["education"])
        for entry in education:
            table = _table(doc, [122, 48])
            row = table.add_row()
            _no_split(row)
            _text_cell(row.cells[0], entry.get("degree", ""), size=9.3, bold=True, color=NAVY)
            _text_cell(row.cells[1], entry.get("dates", ""), size=8.9, color=MUTED,
                       align=WD_ALIGN_PARAGRAPH.RIGHT)
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            _font(p.add_run(entry.get("school", "")), size=9.2, color=MUTED)

    if person.get("languages"):
        _section(doc, HEADINGS["languages"])
        p = doc.add_paragraph()
        _font(p.add_run("  ·  ".join(person["languages"])), size=9.4)

    doc.save(path)


def build_letter(job: dict, profile: dict, body: str, path: Path, lang: str = "de"):
    doc = _doc(kind="letter")
    person = profile.get("person") or {}
    _identity(doc, person)
    de = lang == "de"

    date_line = doc.add_paragraph()
    date_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date_line.paragraph_format.space_before = Pt(4)
    date_line.paragraph_format.space_after = Pt(17)
    _font(date_line.add_run(date.today().strftime("%d.%m.%Y" if de else "%d %B %Y")),
          size=9.5, color=MUTED)

    company = doc.add_paragraph()
    company.paragraph_format.space_after = Pt(7)
    _font(company.add_run(job.get("company") or ""), size=10, color=INK)

    subject = doc.add_paragraph()
    subject.paragraph_format.space_before = Pt(12)
    subject.paragraph_format.space_after = Pt(16)
    subject.paragraph_format.keep_with_next = True
    _font(subject.add_run(("Bewerbung als " if de else "Application for ") + (job.get("title") or "")),
          size=11.4, bold=True, color=NAVY)

    opening = doc.add_paragraph()
    opening.paragraph_format.space_after = Pt(12)
    _font(opening.add_run("Sehr geehrte Damen und Herren," if de else "Dear Hiring Team,"), size=10.5)

    # Each newline from the LLM is a paragraph, as in Stellenpilot's existing renderer.
    for part in body.split("\n"):
        if not part.strip():
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(11)
        p.paragraph_format.keep_together = True
        _font(p.add_run(part.strip()), size=10.65, color=INK)

    closer = doc.add_paragraph()
    closer.paragraph_format.space_before = Pt(13)
    closer.paragraph_format.space_after = Pt(10)
    _font(closer.add_run("Mit freundlichen Grüßen" if de else "Kind regards,"), size=10.5)

    p = doc.add_paragraph()
    _font(p.add_run(person.get("name", "")), size=10.5, bold=True, color=NAVY)
    doc.save(path)
