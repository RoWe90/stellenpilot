"""Truthful CV tailoring, cover letters and a short report per job.

Rules that keep it honest:
* Bullets are copied verbatim from profile.toml. Tailoring only reorders and selects them.
* Generated text (summary, cover letter) passes a truth guard: numbers and known tools must
  appear in your real experience. A tool you do not have yet is allowed only in an honest
  sentence: "I have not used X yet" or "I would like to learn X" for a tool the ad names, and
  "I am learning X" only for a tool in your Learning list. A sentence that also claims use never passes.
* "Covered" and "missing" requirements are always computed by the program, never taken from the LLM.
Limits of the guard: it knows the tools in score.VOCAB and digits. It cannot judge spelled out
numbers ("twelve") or vague claims, so read every document before sending it.
Method credit: inspired by varunr89/resume-tailoring-skill (MIT), see NOTICE.md.
"""
from __future__ import annotations

import re
import shutil
import subprocess
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor

from .score import RE_CACHE, _mentioned, learning_items, profile_text
from .llm import job_ad

# ---------------------------------------------------------------- truth guard

_NUM = re.compile(r"\d+(?:[.,]\d+)?(?:\s?%|[kKmM]\b)?")
# A tool you do not have may appear only in an honest sentence, and which honest sentences are allowed
# depends on the tool:
# * _GAP: you say you have NOT used it yet, or would like to learn it. True for any tool the ad names.
# * _STUDY: you say you ARE learning it now. True only for tools in your Learning list.
# A sentence that also says you used, built or ran something never counts (see _USED), so
# "I used Snowflake daily, a great learning experience" is rejected.
# "machine learning" and "deep learning" do not count; neither do loose words like "not yet" or "vertiefen".
_GAP = re.compile(r"\b(eager to learn|keen to learn|want to learn|would (?:like|love) to learn|glad to learn|happy to learn|looking forward to learning"
                  r"|not used \w+ yet|(?:haven't|have not|have never|never) (?:yet )?(?:used|worked with)|no (?:professional |practical |hands on )?experience (?:yet )?(?:with|in)|new to me"
                  r"|lerne\w*\s+(?:\w+\s+){0,3}gerne?\b|einzuarbeiten|einarbeiten|aneignen|anzueignen"
                  r"|noch keine (?:praktische |berufliche )?erfahrung|noch nicht (?:gearbeitet|eingesetzt|genutzt|verwendet))\b", re.I)
_STUDY = re.compile(r"(?<!machine )(?<!deep )(?<!reinforcement )\b(learning|currently studying|still building"
                    r"|lerne\w*|arbeite\w*\s+(?:ich\s+)?mich\b[^.]{0,40}\bein|einarbeitung)\b", re.I)
_USED = re.compile(r"\b(used|using|built|building|worked|working with|developed|deployed|implemented|designed|maintained|managed"
                   r"|created|wrote|written|ran|running|operated|migrated|automated|delivered|production|professionally|daily"
                   r"|years? of|eingesetzt|genutzt|verwendet|entwickelt|gebaut|umgesetzt|gearbeitet|betreut|einsetze|nutze)\b", re.I)
_SENTENCE = re.compile(r"[^.!?\n]+[.!?]?")


def _honest_kind(sentence: str) -> str:
    """'gap', 'study' or '' for a sentence that mentions a tool you do not have."""
    gap = bool(_GAP.search(sentence))
    rest = _GAP.sub(" ", sentence)
    study = bool(_STUDY.search(rest))
    if _USED.search(_STUDY.sub(" ", rest)):
        return ""
    return "gap" if gap else "study" if study else ""


def _num_tokens(text: str) -> set[str]:
    return {re.sub(r"\s", "", n).lower() for n in _NUM.findall(text or "")}


def unsupported_claims(text: str, profile: dict, job: dict | None = None) -> list[str]:
    """Numbers or tools in `text` that are not backed by the profile. Empty list means safe."""
    text = text or ""
    facts = profile_text(profile, include_learning=False)
    job_head = f"{(job or {}).get('title', '')} {(job or {}).get('company', '')}"
    allowed_nums = _num_tokens(facts) | _num_tokens(job_head)
    bad = [n for n in sorted(_num_tokens(text)) if n not in allowed_nums]
    # Language levels: a letter must never claim more German or English than the profile states.
    from .text import LEVELS, claimed_english_level, profile_language_level, required_german_level
    for lang, claimed in (("German", required_german_level(text)), ("English", claimed_english_level(text))):
        mine = profile_language_level(profile, lang.lower())
        if claimed and (not mine or LEVELS.index(claimed) > LEVELS.index(mine)):
            bad.append(f"{lang} {claimed} (your profile: {mine or 'not stated'})")
    have = _mentioned(facts)
    studying = _mentioned(" ".join(learning_items(profile)))
    in_ad = _mentioned(job.get("title", "") + "\n" + job.get("description", "")) if job else set()
    for sentence in _SENTENCE.findall(text):
        unknown = _mentioned(sentence) - have
        if not unknown:
            continue
        kind = _honest_kind(sentence)
        allowed = studying | in_ad if kind == "gap" else studying if kind == "study" else set()
        bad += [t for t in unknown if t not in allowed]
    return list(dict.fromkeys(bad))


# ---------------------------------------------------------------- planning

PLAN_SYSTEM = """You tailor a CV to a job ad WITHOUT inventing anything.
You may only: choose a headline from the given list, write a 2 to 3 sentence summary that uses ONLY facts from the candidate profile, and reorder or drop items.
Skills under "Currently learning" are NOT experience; never present them as used.
For German job ads, write the CV summary in German as impersonal CV prose (e.g.
"Softwareentwickler mit Erfahrung in ..."), not with "er/sie" and not in first person.
Use concrete profile evidence and relevant projects, not generic enthusiasm.
Do not use hyphens or dashes as punctuation in the summary. Never state language levels or fluency (no 'fluent German', no 'C1'). No numbers or tools that are not in the profile.
Return JSON with keys:
headline_index (int),
summary (string),
skill_groups (list of skill group names, most relevant first; keep all groups),
experience (list, same length as the profile's experience list, each a list of bullet indexes to keep in the new order, at least one),
projects (list of project indexes in the new order; drop clearly irrelevant ones, keep at least one)."""


def _indexed_profile(profile: dict) -> str:
    person = profile.get("person") or {}
    out = [f"HEADLINES: {list(enumerate(person.get('headlines') or []))}",
           f"SKILL GROUPS: {list((profile.get('skills') or {}).keys())}"]
    for i, e in enumerate(profile.get("experience") or []):
        out.append(f"EXPERIENCE {i}: {e.get('title', '')} at {e.get('company', '')}")
        out += [f"  bullet {j}: {b}" for j, b in enumerate(e.get("bullets") or [])]
    for i, p in enumerate(profile.get("projects") or []):
        out.append(f"PROJECT {i}: {p.get('name', '')}: " + " ".join(p.get("bullets") or []))
    return "\n".join(out)


def _relevance(text: str, wanted: set[str]) -> int:
    return sum(1 for t in wanted if RE_CACHE[t].search(text))


def _default_summary(profile: dict) -> str:
    return profile.get("summary") or (profile.get("person") or {}).get("summary") or ""


def coverage(job: dict, profile: dict) -> dict:
    """What the ad asks for versus what the profile shows. Always computed, never from the LLM."""
    wanted = _mentioned(job["title"] + "\n" + job["description"])
    have = _mentioned(profile_text(profile, include_learning=False))
    learning = _mentioned(" ".join(learning_items(profile)))
    return {"coverage": int(100 * len(wanted & have) / len(wanted)) if wanted else 0,
            "covered": sorted(wanted & have),
            "learning": sorted((wanted - have) & learning),
            "missing": sorted(wanted - have - learning)}


def keyword_plan(job: dict, profile: dict) -> dict:
    """No-LLM plan: sort everything by how many of the ad's tools it mentions."""
    wanted = _mentioned(job["title"] + "\n" + job["description"])
    heads = (profile.get("person") or {}).get("headlines") or [""]
    skills = profile.get("skills") or {}
    exp = profile.get("experience") or []
    projs = profile.get("projects") or []
    return {
        "headline_index": max(range(len(heads)), key=lambda i: _relevance(heads[i], wanted)),
        "summary": _default_summary(profile),
        "skill_groups": sorted(skills, key=lambda g: -_relevance(" ".join(skills[g]), wanted)),
        "experience": [sorted(range(len(e.get("bullets") or [])), key=lambda j: -_relevance(e["bullets"][j], wanted)) for e in exp],
        "projects": sorted(range(len(projs)), key=lambda i: -_relevance(" ".join(projs[i].get("bullets") or []), wanted)),
    }


def _valid_plan(plan: dict, profile: dict) -> dict:
    """Clamp indexes so a sloppy LLM answer can never reference facts that do not exist."""
    heads = (profile.get("person") or {}).get("headlines") or [""]
    exp = profile.get("experience") or []
    projs = profile.get("projects") or []
    groups = list(profile.get("skills") or {})
    hi = plan.get("headline_index", 0)
    plan["headline_index"] = hi if isinstance(hi, int) and 0 <= hi < len(heads) else 0
    raw_groups = plan.get("skill_groups") if isinstance(plan.get("skill_groups"), list) else []
    order = [g for g in raw_groups if isinstance(g, str) and g in groups]
    plan["skill_groups"] = list(dict.fromkeys(order)) + [g for g in groups if g not in order]
    exp_plan = plan.get("experience") if isinstance(plan.get("experience"), list) else []
    fixed = []
    for i, e in enumerate(exp):
        n = len(e.get("bullets") or [])
        idx = exp_plan[i] if i < len(exp_plan) and isinstance(exp_plan[i], list) else []
        idx = list(dict.fromkeys(j for j in idx if isinstance(j, int) and 0 <= j < n))
        fixed.append(idx or list(range(n)))
    plan["experience"] = fixed
    raw_p = plan.get("projects") if isinstance(plan.get("projects"), list) else []
    pidx = list(dict.fromkeys(i for i in raw_p if isinstance(i, int) and 0 <= i < len(projs)))
    plan["projects"] = pidx or list(range(len(projs)))
    plan["summary"] = plan["summary"] if isinstance(plan.get("summary"), str) else ""
    return plan


def make_plan(job: dict, profile: dict, llm=None) -> dict:
    plan = keyword_plan(job, profile)
    if llm and llm.enabled:
        user = f"CANDIDATE PROFILE\n{profile_text(profile)}\n\nINDEXED ITEMS\n{_indexed_profile(profile)}\n\nJOB AD\n{job['title']} at {job['company']}\n{job_ad(job['description'], llm.max_chars)}"
        try:
            got = llm.complete_json(PLAN_SYSTEM, user, max_tokens=1200)
            plan.update({k: got[k] for k in ("headline_index", "summary", "skill_groups", "experience", "projects") if k in got})
        except Exception as e:  # fall back to the keyword plan, never block the run
            plan["note"] = f"LLM plan failed, used keyword plan: {e}"
    plan = _valid_plan(plan, profile)
    bad = unsupported_claims(plan["summary"], profile, job)
    if bad:
        plan["summary"] = _default_summary(profile)
        plan["guard"] = f"Generated summary mentioned {bad}, not in your experience, so the default summary was used."
    plan.update(coverage(job, profile))
    return plan



class ProfileQualityError(ValueError):
    """An imported CV mixed job headings and job descriptions; manual correction is required."""


def profile_quality_issues(profile: dict) -> list[str]:
    """Detect obvious import corruption without rewriting someone's career history.

    This is deliberately conservative: a long but legitimate bullet must survive,
    while a job title/company full of bullet glyphs or narrative is not a valid heading.
    """
    issues: list[str] = []
    for idx, role in enumerate(profile.get("experience") or [], 1):
        for field, label in (("title", "Titel"), ("company", "Arbeitgeber")):
            value = str(role.get(field) or "")
            if any(mark in value for mark in ("•", "●", "▪", "", "\n", "\r")):
                issues.append(f"Berufsstation {idx}: {label} enthält Aufzählungen oder Zeilenumbrüche")
            elif len(value) > 135:
                issues.append(f"Berufsstation {idx}: {label} ist ungewöhnlich lang ({len(value)} Zeichen)")
        # PDF text extraction can shift the entire job description into the employer field
        # without leaving bullet glyphs. Catch the common long "Company, city, Development ..."
        # pattern, but keep ordinary short employer/location combinations valid.
        company = str(role.get("company") or "")
        if (len(company) > 70 and company.count(",") >= 2
                and re.search(r"(?i),\\s*(?:Entwicklung|Implementierung|Automatisierung|"
                              r"Werkzeugentwicklung|Daten(?:vorverarbeitung|transformation))\\b", company)):
            issues.append(f"Berufsstation {idx}: Arbeitgeber enthält offenbar Tätigkeitsbeschreibungen")
        bullets = [str(b).strip() for b in (role.get("bullets") or []) if str(b).strip()]
        normalized = [re.sub(r"\\s+", " ", b).casefold() for b in bullets]
        if len(normalized) != len(set(normalized)):
            issues.append(f"Berufsstation {idx}: doppelte Tätigkeitsbeschreibungen")
        for bullet in bullets:
            if any(glyph in bullet for glyph in ("•", "●", "▪", "")):
                issues.append(f"Berufsstation {idx}: Stichpunkt enthält PDF-Aufzählungszeichen")
                break
            if bullet.count("(") != bullet.count(")"):
                issues.append(f"Berufsstation {idx}: Stichpunkt enthält zerrissene Klammerausdrücke")
                break
            if re.match(r"(?i)^(?:mit dem Thema|und (?:Implementierung|Hindernisvermeidung)|"
                        r"zum ferngesteuerten|des Backends|Pipelines, CI/CD)\\b", bullet):
                issues.append(f"Berufsstation {idx}: Stichpunkt ist vermutlich ein Satzfragment")
                break
    for idx, education in enumerate(profile.get("education") or [], 1):
        degree = str(education.get("degree") or "")
        school = str(education.get("school") or "")
        if (degree.count("(") != degree.count(")") or school.count("(") != school.count(")")
                or school.lower().lstrip().startswith("umorientierung")):
            issues.append(f"Ausbildung {idx}: Abschluss und Schule wurden beim Import offenbar falsch getrennt")
    return issues


# ---------------------------------------------------------------- documents

def _para(doc, text: str, size: float | None = None, bold: bool = False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
    return p


# CV layouts. Layout only: content always comes from the plan, so the truth guard covers every style.
EN_HEADS = {"summary": "Professional Summary", "skills": "Skills", "experience": "Experience",
            "projects": "Projects", "education": "Education", "languages": "Languages"}
DE_HEADS = {"summary": "Profil", "skills": "Kenntnisse", "experience": "Berufserfahrung",
            "projects": "Projekte", "education": "Ausbildung", "languages": "Sprachen"}
CV_STYLES = {
    "classic": {"size": 10.5, "margin": 50, "heads": EN_HEADS, "dates_first": False},
    "compact": {"size": 9.5, "margin": 36, "heads": EN_HEADS, "dates_first": False},
    "lebenslauf": {"size": 10, "margin": 43, "heads": DE_HEADS, "dates_first": True},
}


def _heading(doc, text):
    p = _para(doc, text.upper(), 10.5, bold=True)
    p.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True


def _base_doc(size: float = 10.5, margin: int = 50) -> Document:
    """DIN A4 on all platforms, with quiet, consistent paragraph typography."""
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Carlito"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0x20, 0x29, 0x36)
    st.paragraph_format.line_spacing = 1.08
    st.paragraph_format.space_after = Pt(3)
    bullets = doc.styles["List Bullet"]
    bullets.font.name = "Carlito"
    bullets.font.size = Pt(size)
    bullets.paragraph_format.space_after = Pt(2)
    bullets.paragraph_format.left_indent = Pt(16)
    bullets.paragraph_format.first_line_indent = Pt(-7)
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = section.right_margin = Pt(margin)
        section.top_margin = section.bottom_margin = Pt(max(margin - 7, 31))
    return doc


def _contact(person: dict, links: bool = True) -> str:
    parts = [person.get("location"), person.get("email"), person.get("phone")]
    if links:
        parts += person.get("links") or []
    return " | ".join(x for x in parts if x)


def build_cv(profile: dict, plan: dict, path: Path, style: str = "classic") -> None:
    st = CV_STYLES.get(style, CV_STYLES["classic"])
    h = st["heads"]
    person = profile.get("person") or {}
    doc = _base_doc(st["size"], st["margin"])
    _para(doc, person.get("name", ""), 20, bold=True)
    heads = person.get("headlines") or []
    if heads:
        _para(doc, heads[plan["headline_index"]], 11.5)
    _para(doc, _contact(person), 9.5)

    if plan["summary"]:
        _heading(doc, h["summary"])
        _para(doc, plan["summary"])

    skills = profile.get("skills") or {}
    if skills:
        _heading(doc, h["skills"])
        for g in plan["skill_groups"]:
            p = doc.add_paragraph()
            p.add_run(f"{g}: ").bold = True
            p.add_run(", ".join(skills[g]))
            p.paragraph_format.space_after = Pt(0)

    exp = profile.get("experience") or []
    if exp:
        _heading(doc, h["experience"])
        for e, idx in zip(exp, plan["experience"]):
            p = doc.add_paragraph()
            if st["dates_first"] and e.get("dates"):
                p.add_run(f"{e['dates']}    ").italic = True
            p.add_run(", ".join(x for x in (e.get("title"), e.get("company")) if x)).bold = True
            if e.get("dates") and not st["dates_first"]:
                p.add_run(f"    {e['dates']}").italic = True
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = bool(idx)
            p.paragraph_format.keep_together = True
            for j in idx:
                bullet = doc.add_paragraph(e["bullets"][j], style="List Bullet")
                bullet.paragraph_format.keep_together = True

    projs = profile.get("projects") or []
    if projs:
        _heading(doc, h["projects"])
        for i in plan["projects"]:
            pr = projs[i]
            p = doc.add_paragraph()
            p.add_run(pr.get("name", "")).bold = True
            if pr.get("link"):
                p.add_run(f"    {pr['link']}").italic = True
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = bool(pr.get("bullets"))
            for b in pr.get("bullets") or []:
                bullet = doc.add_paragraph(b, style="List Bullet")
                bullet.paragraph_format.keep_together = True

    if profile.get("education"):
        _heading(doc, h["education"])
        for ed in profile["education"]:
            p = doc.add_paragraph()
            p.add_run(ed.get("degree", "")).bold = True
            rest = ", ".join(x for x in (ed.get("school"), ed.get("dates")) if x)
            if rest:
                p.add_run(f", {rest}")
            p.paragraph_format.space_after = Pt(0)
    if person.get("languages"):
        _heading(doc, h["languages"])
        _para(doc, ", ".join(person["languages"]))
    doc.save(path)


LETTER_SYSTEM = """Write the BODY of a concise, evidence-led job application letter in English.
Use 3 short paragraphs, 140 to 220 words. No greeting, subject line or signature.
Paragraph 1: a concrete reason this specific role fits the candidate, tied to work already done.
Paragraph 2: one or two RELEVANT work examples from the profile, describing actions, tools and
results actually documented. Paragraph 3: connect these examples to the job's priorities and
close with a confident invitation to discuss.
Use ONLY verified profile facts; do not invent years, achievements, languages or expertise.
Skills marked Currently learning are NOT experience. Do not volunteer missing skills or
declare that the candidate lacks tools merely because their names do not occur in the profile.
Never claim experience with something absent from the profile. Avoid clichés such as
'I hereby apply', 'I am very passionate', and empty adjectives.
Write naturally in the first person. No fabricated enthusiasm or personal knowledge of the firm.
Return JSON only: {"body": "..."}"""


LETTER_SYSTEM_DE = """Verfasse ausschließlich den HAUPTTEIL eines überzeugenden deutschen Bewerbungsanschreibens.
Drei kurze Absätze, etwa 140 bis 220 Wörter; keine Anrede, Betreffzeile oder Grußformel.
Absatz 1: konkreter Bezug zur ausgeschriebenen Tätigkeit und zu belegbarer Erfahrung.
Absatz 2: ein bis zwei passende, überprüfbare Beispiele aus der Berufspraxis bzw. Projekten.
Absatz 3: deren Nutzen für die konkrete Stelle und ein sachlicher, selbstbewusster Abschluss.
Verwende nur nachweisbare Profilfakten. Keine erfundenen Erfolge, Zeitangaben, Tools oder Kenntnisse.
'Learning' ist keine Berufserfahrung. Liste keine vermeintlichen Defizite auf und behaupte
nicht, dass eine Technologie unbekannt sei, nur weil sie nicht im Profil erwähnt wird.
Keine Floskeln wie 'hiermit bewerbe ich mich', 'reizt mich sehr', 'ehrlich gesagt'.
Schreibe in der ersten Person, aktiv, präzise und natürlich. Keine Sprachlevel erwähnen.
Keine Gedankenstriche als Satzzeichen. Gib ausschließlich JSON zurück: {"body": "..."}"""


def _gap_sentence(plan: dict) -> str:
    """Learning list tools are 'currently learning'; other gaps are 'not used yet'. Never more."""
    learning, missing = plan.get("learning", [])[:2], plan.get("missing", [])[:2]
    parts = []
    if learning:
        parts.append(f"I am currently learning {', '.join(learning)}.")
    if missing and len(learning) < 2:
        parts.append(f"I have not worked with {', '.join(missing[:2 - len(learning)])} yet and would be glad to learn it on the job.")
    return " ".join(parts)


def _gap_sentence_de(plan: dict) -> str:
    learning, missing = plan.get("learning", [])[:2], plan.get("missing", [])[:2]
    parts = []
    if learning:
        parts.append(f"In {', '.join(learning)} arbeite ich mich gerade ein.")
    if missing and len(learning) < 2:
        parts.append(f"Mit {', '.join(missing[:2 - len(learning)])} habe ich noch nicht gearbeitet, lerne es aber gern im Job.")
    return " ".join(parts)


def template_letter_de(job: dict, plan: dict) -> str:
    covered = ", ".join(plan.get("covered", [])[:4])
    gaps = plan.get("learning", []) + plan.get("missing", [])
    p1 = f"hiermit bewerbe ich mich auf die Stelle {job['title']} bei {job['company']}."
    p2 = (f"In Ihrer Ausschreibung nennen Sie {covered}. Damit habe ich in meiner bisherigen Arbeit und in Projekten gearbeitet. " if covered
          else "Die Aufgaben passen gut zu dem, was ich bisher in meiner Arbeit und in Projekten gemacht habe. ")
    p3 = "Gern erläutere ich Ihnen im Gespräch, wie ich meine Erfahrungen in die ausgeschriebene Tätigkeit einbringen kann."
    return "\n\n".join([p1, p2.strip(), p3])


def template_letter(job: dict, plan: dict) -> str:
    covered = ", ".join(plan.get("covered", [])[:4])
    gaps = plan.get("learning", []) + plan.get("missing", [])
    para1 = f"I am applying for the {job['title']} role at {job['company']}. {plan['summary']}".strip()
    para2 = (f"Your ad mentions {covered}, which I have used in my work and projects. " if covered else
             "The role fits the work I have done in my previous positions and projects. ")
    para3 = "I would welcome the opportunity to discuss how my experience relates to this role."
    return "\n\n".join([para1, para2.strip(), para3])


def letter_language(job: dict, setting: str = "auto", llm=None, my_german: str = "A2") -> str:
    """'auto' writes German only for German ads AND when your German is B2 or better. A polished
    German letter from someone at A2 would suggest a level they do not have."""
    from .text import LEVELS, ad_language
    if setting in ("en", "de"):
        return setting
    fluent_enough = my_german.upper() in LEVELS and LEVELS.index(my_german.upper()) >= LEVELS.index("B2")
    german_ad = ad_language(job["title"] + "\n" + job["description"]) == "de"
    return "de" if german_ad and fluent_enough and llm and llm.enabled else "en"


def letter_body(job: dict, profile: dict, plan: dict, llm=None, lang: str = "en") -> tuple[str, str]:
    """Returns (body, note). Falls back to a plain template when the guard rejects the draft."""
    template = template_letter_de(job, plan) if lang == "de" else template_letter(job, plan)
    if not (llm and llm.enabled):
        return template, "Template letter (no LLM configured)."
    name = (profile.get("person") or {}).get("name", "")
    user = f"CANDIDATE PROFILE\n{profile_text(profile)}\n\nJOB\n{job['title']} at {job['company']}\n{job_ad(job['description'], llm.max_chars)}\n\nCandidate name: {name}"
    bad: list[str] = []
    for _ in range(2):
        try:
            body = str(llm.complete_json(LETTER_SYSTEM_DE if lang == "de" else LETTER_SYSTEM, user, max_tokens=900).get("body") or "")
        except Exception as e:
            return template, f"LLM letter failed ({e}), template used."
        bad = unsupported_claims(body, profile, job)
        if body and not bad:
            return body, ""
        user += f"\n\nYour previous draft mentioned {bad}, which are not backed by the profile. Remove them. Only if the ad asks for one, you may say plainly that the candidate has not used it yet."
    return template, f"LLM drafts mentioned things not in your experience ({bad}), so a template letter was used."


def build_letter(job: dict, profile: dict, body: str, path: Path, lang: str = "en") -> None:
    person = profile.get("person") or {}
    doc = _base_doc(10.5, 49)
    _para(doc, person.get("name", ""), 18, bold=True)
    _para(doc, _contact(person, links=False), 9.5)
    de = lang == "de"
    dp = _para(doc, date.today().strftime("%d.%m.%Y" if de else "%d %B %Y"))
    dp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    dp.paragraph_format.space_before = Pt(12)
    company = _para(doc, job["company"])
    company.paragraph_format.space_before = Pt(13)
    subject = _para(doc, f"{'Bewerbung als' if de else 'Re:'} {job['title']}", 11, bold=True)
    subject.paragraph_format.space_before = Pt(12)
    subject.paragraph_format.space_after = Pt(14)
    salutation = _para(doc, "Sehr geehrte Damen und Herren," if de else "Dear Hiring Team,")
    salutation.paragraph_format.space_after = Pt(10)
    for para in [p.strip() for p in body.split("\n") if p.strip()]:
        paragraph = _para(doc, para)
        paragraph.paragraph_format.space_after = Pt(8)
    ending = _para(doc, "Mit freundlichen Grüßen" if de else "Kind regards,")
    ending.paragraph_format.space_before = Pt(8)
    _para(doc, person.get("name", ""))
    doc.save(path)


MAC_SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
WIN_SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"


def find_soffice() -> str | None:
    """LibreOffice is often not on PATH on macOS and Windows, so check the default install places."""
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for p in (MAC_SOFFICE, WIN_SOFFICE):
        if Path(p).exists():
            return p
    return None


def to_pdf(paths: list[Path]) -> list[Path]:
    """PDF copies next to the .docx files. Never fails the tailoring: no LibreOffice, a crash or a
    timeout just means no PDF. Each call uses its own LibreOffice profile, because a LibreOffice
    window that is already open otherwise swallows the request and no PDF appears."""
    import tempfile
    exe = find_soffice()
    if not exe:
        return []
    outdir = paths[0].parent
    with tempfile.TemporaryDirectory(prefix="stellenpilot-lo-", ignore_cleanup_errors=True) as profile:
        try:
            subprocess.run([exe, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", "pdf",
                            "--outdir", str(outdir), *map(str, paths)],
                           check=False, capture_output=True, timeout=180)
        except (subprocess.TimeoutExpired, OSError):
            pass
    return [p.with_suffix(".pdf") for p in paths if p.with_suffix(".pdf").exists()]


_UMLAUT = str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue", "Ä": "Ae", "Ö": "Oe", "Ü": "Ue", "ß": "ss"})


def slug(s: str, n: int = 40) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", (s or "").translate(_UMLAUT)).strip("_")[:n].strip("_") or "job"


def tailor_job(job: dict, profile: dict, out_root: Path, llm=None, make_pdf: bool = True,
               letter_setting: str = "auto", my_german: str = "A2", cv_style: str = "classic") -> dict:
    issues = profile_quality_issues(profile)
    if issues:
        raise ProfileQualityError("Bitte profile.toml prüfen: " + "; ".join(issues[:6]))
    person = profile.get("person") or {}
    # Company plus title, so two roles at one company never overwrite each other.
    folder = out_root / date.today().isoformat() / f"{slug(job['company'], 30)}__{slug(job['title'], 40)}"
    folder.mkdir(parents=True, exist_ok=True)
    from . import progress
    steps = 4 if make_pdf else 3
    progress.stage("Choosing the experience, projects and skills that fit this job", expect=700, step=1, steps=steps)
    plan = make_plan(job, profile, llm)
    who = slug(person.get("name", "CV"))
    cv = folder / f"{who}_CV_{slug(job['company'], 30)}.docx"
    build_cv(profile, plan, cv, cv_style)
    lang = letter_language(job, letter_setting, llm, my_german)
    progress.stage("Writing the cover letter", expect=1500, step=2, steps=steps)
    body, note = letter_body(job, profile, plan, llm, lang)
    progress.stage("Making the Word files", step=3, steps=steps)
    from .text import LEVELS, normalize_level
    level_note = ""
    if lang == "de" and LEVELS.index(normalize_level(my_german) or "A2") < LEVELS.index("B2"):
        level_note = (f"This letter is in German, but your German is {normalize_level(my_german) or 'A2'}. A polished German "
                      "letter suggests a higher level than you have; consider letter_language = \"en\" or say your level in the letter.")
    letter = folder / f"{who}_{'Anschreiben' if lang == 'de' else 'Cover_Letter'}_{slug(job['company'], 30)}.docx"
    build_letter(job, profile, body, letter, lang)
    report = folder / "report.md"
    report.write_text("\n".join([
        f"# {job['title']} at {job['company']}",
        f"Match: {job.get('match')}  |  Tools the ad names that your experience shows: {plan['coverage']}%",
        f"Link: {job.get('url', '')}", "",
        "## Covered by your experience", *[f"* {c}" for c in plan["covered"]], "",
        "## You are learning these", *[f"* {c}" for c in plan["learning"]], "",
        "## Missing (say so honestly if asked)", *[f"* {m}" for m in plan["missing"]], "",
        "## Notes", "* Read both documents before sending. The truth guard checks tools and numbers, not every claim.",
        *[f"* {n}" for n in (level_note, plan.get("note"), plan.get("guard"), note) if n],
    ]) + "\n", encoding="utf-8")
    if make_pdf:
        progress.stage("Making the PDF files", step=4, steps=steps)
    pdfs = to_pdf([cv, letter]) if make_pdf else []
    return {"folder": str(folder), "files": [str(p) for p in (cv, letter, report, *pdfs)],
            "coverage": plan["coverage"], "cv_style": cv_style if cv_style in CV_STYLES else "classic"}
