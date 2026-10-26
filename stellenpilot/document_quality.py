"""Offline, privacy-preserving review of generated application documents.

Checks useful invariants without pretending to certify visual aesthetics or
the factual accuracy of personal claims. No documents are sent to a service.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

DE_FIXES = (
    (re.compile(r"\bPython[ -]Datenverarbeitung\b", re.I), "Python-Datenverarbeitung"),
    (re.compile(r"\bETL[ -]Prozessen\b", re.I), "ETL-Prozessen"),
    (re.compile(r"\bKI[ -]Agenten[ -]Lösung\b", re.I), "KI-Agenten-Lösung"),
    (re.compile(r"\bKI[ -]Agenten[ -]basiert", re.I), "KI-agentenbasiert"),
    (re.compile(r"\bCI[ -]CD\b", re.I), "CI/CD"),
)


def polish_generated_german(text: str) -> str:
    """Conservative compound spelling fixes, ONLY for LLM-generated prose.

    Never use this on profile facts, employer names or original CV bullet text.
    """
    for pattern, replacement in DE_FIXES:
        text = pattern.sub(replacement, text or "")
    return text


def letter_quality_issues(text: str, *, lang: str) -> list[str]:
    """Formatting/style guidance for a single retry, not a factual truth check."""
    if not (text or "").strip():
        return ["empty draft"]
    words = re.findall(r"\b[\wÄÖÜäöüß]+\b", text)
    problems = []
    if len(words) < 85:
        problems.append("too short; use concrete, documented job examples")
    if len(words) > 270:
        problems.append("too long; focus on the two most relevant examples")
    paragraphs = [p for p in re.split(r"\n\s*\n", text.strip()) if p.strip()]
    if len(paragraphs) != 3:
        problems.append("not three coherent paragraphs")
    if re.search(r"(?i)\b(hiermit bewerbe ich mich|ehrlich gesagt|I hereby apply)\b", text):
        problems.append("generic or defensive phrasing")
    if lang == "de" and re.search(r"(?i)\b(Sehr geehrte|Mit freundlichen Grüßen)\b", text):
        problems.append("contains greeting/sign-off already provided by document renderer")
    return problems



# A small set of material action claims that mere tool presence cannot support.
# This is intentionally not presented as a comprehensive semantic fact checker.
_ACCEPTANCE = re.compile(r"\b(?:Akzeptanzkriterien|acceptance criteria)\b", re.I)
_DEPLOYMENT = re.compile(r"\b(?:Auslieferung|Deployment|deploy(?:ed|ment)?|"
                         r"Produktivbetrieb|produktiv ausgerollt)\b", re.I)
_CI_TOOLS = re.compile(r"\b(?:Azure Pipelines?|CI/CD|Docker|GitHub Actions|"
                       r"Jenkins|GitLab CI)\b", re.I)
_PROSPECTIVE = re.compile(
    r"\b(?:möchte|würde|künftig|will ich lernen|interessiere mich|"
    r"would like|hope to|aim to|look forward to)\b", re.I
)
_OWN_EXPERIENCE = re.compile(
    r"\b(?:ich|mir|meine|meiner|mein|habe|nutze|setze|"
    r"my|i have|i use|i built|i delivered)\b", re.I
)


def unsupported_action_claims(text: str, profile_text: str) -> list[str]:
    """Identify specific *experience* claims unsupported by the candidate facts.

    The check is deliberately conservative. A mention of a prospective role
    requirement is NOT the same thing as a claim of past experience. These
    rules catch the previously observed acceptance-criteria and deployment
    inferences, which technology-name/number-only guards miss.
    """
    issues = []
    facts = profile_text or ""
    sentences = [piece.strip() for piece in re.split(r"(?<=[.!?])\s+|\n+", text or "")
                 if piece.strip()]
    for sentence in sentences:
        prospective = bool(_PROSPECTIVE.search(sentence))
        first_person = bool(_OWN_EXPERIENCE.search(sentence))
        if (_ACCEPTANCE.search(sentence) and first_person and not prospective
                and not _ACCEPTANCE.search(facts)):
            issues.append("unsupported_experience_acceptance_criteria")
        if (_DEPLOYMENT.search(sentence) and _CI_TOOLS.search(sentence)
                and first_person and not prospective
                and not any(_DEPLOYMENT.search(line) and _CI_TOOLS.search(line)
                            for line in facts.splitlines())):
            issues.append("unsupported_experience_deployment")
    return list(dict.fromkeys(issues))

def inspect_documents(cv: Path, letter: Path, pdfs: list[Path] | None = None,
                      expected_bullets: list[str] | None = None,
                      require_pdf: bool = True) -> dict:
    """Return serializable issues and counts: errors mean broken artifact.

    PDF pagination and text density are advisory (warnings). PDF pages are
    read one at a time; no user data are logged in the diagnostic details.
    """
    issues: list[dict] = []
    stats: dict = {}
    pdfs = list(pdfs or [])
    for kind, path in (("CV", cv), ("Anschreiben", letter)):
        if not path.is_file() or path.stat().st_size == 0:
            issues.append({"severity": "error", "code": "missing_docx", "detail": kind})
            continue
        try:
            doc = Document(path)
            width, height = doc.sections[0].page_width.mm, doc.sections[0].page_height.mm
            if not 209 <= width <= 211 or not 296 <= height <= 298:
                issues.append({"severity": "error", "code": "not_a4_docx", "detail": kind})
            text = "\n".join(p.text for p in doc.paragraphs)
            if not text.strip():
                issues.append({"severity": "error", "code": "empty_docx", "detail": kind})
            if kind == "CV":
                missed = [b for b in expected_bullets or [] if b not in text]
                if missed:
                    issues.append({"severity": "error", "code": "missing_source_bullets",
                                   "detail": f"{len(missed)} source bullet(s) absent"})
                if doc.tables:
                    issues.append({"severity": "warning", "code": "cv_tables",
                                   "detail": "Table-based content may disrupt reading order"})
        except Exception as exc:
            issues.append({"severity": "error", "code": "unreadable_docx",
                           "detail": f"{kind}: {type(exc).__name__}"})
    if require_pdf and len(pdfs) < 2:
        issues.append({"severity": "warning", "code": "pdf_not_created",
                       "detail": "PDF export unavailable or incomplete; DOCX is still usable"})
    for pdf in pdfs:
        kind = "CV" if pdf.stem == cv.stem else "Anschreiben" if pdf.stem == letter.stem else "Unbekannt"
        if kind == "Unbekannt":
            issues.append({"severity": "warning", "code": "unexpected_pdf", "detail": pdf.name})
            continue
        try:
            reader = PdfReader(str(pdf))
            pages = len(reader.pages)
            stats["cv_pages" if kind == "CV" else "letter_pages"] = pages
            if pages < 1:
                issues.append({"severity": "error", "code": "empty_pdf", "detail": kind})
                continue
            if pages > (2 if kind == "CV" else 1):
                issues.append({"severity": "warning", "code": "too_many_pages",
                               "detail": f"{kind}: {pages} pages"})
            for i, page in enumerate(reader.pages):
                width, height = float(page.mediabox.width), float(page.mediabox.height)
                if not (591 <= width <= 597 and 839 <= height <= 845):
                    issues.append({"severity": "error", "code": "not_a4_pdf",
                                   "detail": f"{kind}: page {i + 1}"})
                content = page.extract_text() or ""
                if len(content.strip()) < 180:
                    issues.append({"severity": "warning", "code": "sparse_page",
                                   "detail": f"{kind}: page {i + 1} has very little text"})
                if re.search(r"(?:KOMPETENZEN|BERUFSERFAHRUNG|AUSBILDUNG|SPRACHEN)\s*$",
                             content.strip()):
                    issues.append({"severity": "warning", "code": "orphaned_heading",
                                   "detail": f"{kind}: page {i + 1}"})
        except Exception as exc:
            issues.append({"severity": "error", "code": "unreadable_pdf",
                           "detail": f"{kind}: {type(exc).__name__}"})
    return {"issues": issues, "stats": stats, "ok": not any(i["severity"] == "error" for i in issues)}
